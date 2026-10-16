import io
import re
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from .config import settings


def parse_document(data: bytes, filename: str) -> list[dict]:
    if not data or len(data) > settings().max_upload_bytes:
        raise ValueError('Upload a non-empty file no larger than 10 MB.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError('Password-protected PDFs are not supported. Upload an unlocked copy.')
        if len(reader.pages) > 200:
            raise ValueError('Limit each document to 200 pages.')
        pages = [{'number': i + 1, 'text': p.extract_text() or ''} for i, p in enumerate(reader.pages)]
    elif suffix == '.docx':
        # Guard compressed documents before python-docx expands the archive.
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 30_000_000:
                raise ValueError('Expanded DOCX is too large.')
        doc = Document(io.BytesIO(data))
        text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        # DOCX has no reliable source page numbers. Viewer uses explicitly labeled sections.
        pages = section_text(text, 'section')
    elif suffix in {'.txt', '.md'}:
        pages = section_text(data.decode('utf-8-sig'), 'section')
    else:
        raise ValueError('Use PDF, DOCX, TXT or Markdown files.')
    total = sum(len(p['text']) for p in pages)
    if total < 80:
        raise ValueError('Too little readable text. Scanned documents need OCR before upload.')
    if total > settings().max_document_chars:
        raise ValueError('Document exceeds 180,000 readable characters. Split the document.')
    for p in pages:
        p.setdefault('kind', 'page')
    return pages


def section_text(text: str, kind='section') -> list[dict]:
    # Preserve exact characters within each section; offsets always refer to this stored text.
    sections = re.split(r'\f', text)
    if len(sections) == 1:
        sections = [text[i:i + 6000] for i in range(0, len(text), 6000)]
    return [{'number': i + 1, 'text': t, 'kind': kind} for i, t in enumerate(t for t in sections if t.strip())]
