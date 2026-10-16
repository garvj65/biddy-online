import asyncio
import io
from datetime import timedelta
from unittest.mock import patch
import httpx
import pytest
from docx import Document
from app.db import Job, Session, Solicitation, now
from app.documents import parse_document, section_text
from app.extraction import analyze, chunks, verify_facts
from app.providers import ProviderError, extract
from app.schemas import RawAnalysis
from app.worker import claim_job, renew_lease, run_once


def raw(quote, page=1, value=None):
    return RawAnalysis.model_validate({'title': 'Example', 'agency': None, 'solicitation_number': None,
        'facts': [{'category': 'deadline', 'label': 'Due', 'value': value or quote, 'quote': quote, 'page': page, 'severity': 'normal'}], 'missing_information': []})


def test_hallucinated_quotes_and_wrong_pages_are_withheld():
    pages = [{'number': 1, 'text': 'Proposals are due Friday at 5 PM ET.'}]
    assert verify_facts(raw('Proposals are due Monday at 5 PM ET.'), pages, 'a') == ([], 1)
    assert verify_facts(raw(pages[0]['text'], page=2), pages, 'a') == ([], 1)
    facts, rejected = verify_facts(raw(pages[0]['text']), pages, 'a')
    assert rejected == 0 and len(facts) == 1


def test_chunking_preserves_pages_and_full_document():
    text = 'x' * 32000
    result = list(chunks([{'number': 7, 'text': text}], budget=14000))
    assert len(result) == 3
    assert all(p['number'] == 7 for group in result for p in group)
    assert sum(len(p['text']) for group in result for p in group) >= len(text)


def test_docx_has_sections_not_invented_pdf_pages():
    doc = Document()
    doc.add_paragraph('EXAMPLE PROCUREMENT')
    doc.add_paragraph('Proposals are due October 30, 2026 at 5 PM Eastern Time. The contractor shall provide Drupal support.')
    buf = io.BytesIO()
    doc.save(buf)
    pages = parse_document(buf.getvalue(), 'solicitation.docx')
    assert pages[0]['kind'] == 'section' and 'Drupal' in pages[0]['text']


def test_readable_pdf_preserves_page_and_text():
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 50 700 Td (EXAMPLE SOLICITATION. Proposals are due October 30, 2026 at 5 PM Eastern Time. The contractor shall provide cloud support.) Tj ET')
    page[NameObject('/Contents')] = stream
    buf = io.BytesIO()
    writer.write(buf)
    pages = parse_document(buf.getvalue(), 'example.pdf')
    assert pages[0]['number'] == 1 and pages[0]['kind'] == 'page'
    assert 'Eastern Time' in pages[0]['text']


def test_empty_text_sections_do_not_break_viewer_navigation():
    pages = section_text('First section\f\fThird section')
    assert [p['number'] for p in pages] == [1, 2]


def test_blank_pdf_cover_does_not_break_analysis():
    result = asyncio.run(analyze([{'number': 1, 'text': ''}, {'number': 2, 'text': 'EXAMPLE PROCUREMENT\nProposals are due October 30, 2026 at 5 PM Eastern Time.'}], 'a'))
    assert result['facts'][0]['citation']['page'] == 2
    assert result['title'] == 'EXAMPLE PROCUREMENT'


def test_missing_fields_remain_unknown():
    result = asyncio.run(analyze([{'number': 1, 'text': 'The contractor shall provide helpdesk support and incident reporting.'}], 'a'))
    assert any('Response deadline' in x for x in result['missing_information'])
    assert not any(f['category'] == 'deadline' for f in result['facts'])
    assert result['recommendation'] == 'review_required'


def test_absent_due_date_is_not_a_deadline_and_timezone_gaps_are_visible():
    result = asyncio.run(analyze([{'number': 1, 'text': 'EXAMPLE\nNo response schedule or closing date is specified in this excerpt.\nQuestions must be submitted by October 20, 2026 at noon ET.'}], 'a'))
    assert len([f for f in result['facts'] if f['category'] == 'deadline']) == 1
    assert any('Response deadline' in m for m in result['missing_information'])
    ambiguous = asyncio.run(analyze([{'number': 1, 'text': 'Proposals are due November 12, 2026 at 5:00 PM.'}], 'a'))
    assert any('timezone' in m for m in ambiguous['missing_information'])


def test_provider_adapter_schema_and_key_not_in_errors(monkeypatch):
    from app.config import settings
    s = settings()
    monkeypatch.setattr(s, 'llm_provider', 'openai')
    monkeypatch.setattr(s, 'llm_api_key', 'secret-not-for-logs')
    seen = []
    def handler(request):
        import json
        payload = json.loads(request.content)
        assert payload['response_format']['json_schema']['strict'] is True
        assert request.headers['authorization'] == 'Bearer secret-not-for-logs'
        seen.append(request)
        return httpx.Response(401, json={'error': 'secret-not-for-logs'})
    async def invoke():
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            with pytest.raises(ProviderError) as exc:
                await extract([{'number': 1, 'text': 'Example procurement'}], client)
            assert 'secret-not-for-logs' not in str(exc.value)
    asyncio.run(invoke())
    assert len(seen) == 1


def test_native_anthropic_and_gemini_adapters(monkeypatch):
    import json
    from app.config import settings
    s = settings()
    monkeypatch.setattr(s, 'llm_api_key', 'test-key')
    obj = raw('Proposals are due Friday at 5 PM ET.').model_dump()
    async def invoke(provider):
        monkeypatch.setattr(s, 'llm_provider', provider)
        def handler(request):
            body = json.loads(request.content)
            if provider == 'anthropic':
                assert body['tool_choice']['name'] == 'extract_solicitation'
                return httpx.Response(200, json={'content': [{'type': 'tool_use', 'input': obj}], 'usage': {'input_tokens': 10}})
            assert body['generationConfig']['responseMimeType'] == 'application/json'
            return httpx.Response(200, json={'candidates': [{'content': {'parts': [{'text': json.dumps(obj)}]}}]})
        async with httpx.AsyncClient(transport=httpx.MockTransport(handler)) as client:
            result, _ = await extract([{'number': 1, 'text': obj['facts'][0]['quote']}], client)
            assert result.facts[0].page == 1
    for provider in ['anthropic', 'gemini']:
        asyncio.run(invoke(provider))


def test_retry_backoff_and_stale_lease_recovery(client, headers, solicitation_text, monkeypatch):
    queued = client.post('/api/solicitations/text', headers=headers, json={'title': 'Example', 'text': solicitation_text}).json()
    async def fail(*args):
        raise ProviderError('Provider unavailable.')
    with patch('app.worker.analyze', fail):
        assert run_once()
    with Session.begin() as db:
        job = db.get(Job, queued['job_id'])
        assert job.attempts == 1 and job.status == 'queued'
        job.available_at = now() - timedelta(seconds=1)
    claimed = claim_job()
    assert claimed and not claim_job()
    assert not renew_lease(claimed[0], 'not-the-owning-token')
    assert renew_lease(claimed[0], claimed[1])
    with Session.begin() as db:
        job = db.get(Job, queued['job_id'])
        job.claimed_at = now() - timedelta(hours=2)
    assert run_once()
    with Session() as db:
        assert db.get(Job, queued['job_id']).status == 'done'
        assert db.get(Solicitation, queued['id']).status == 'ready'
